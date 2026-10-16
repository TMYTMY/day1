"""Load a CSV or Excel file and write its declaration.

Writing ``config/sources/<source_id>.yaml`` is the declaration. A human runs this
command. A web upload must never call it.

Excel workbooks are read with openpyxl (``data_only=True``). Legacy ``.xls`` is
rejected. Formula cells with no cached value are staged as NULL and counted.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import sqlite3
import sys
import zipfile
from datetime import date, datetime, time, timezone
from pathlib import Path

import pandas
import yaml
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from src.config import ConfigError
from src.sources import (
    DEFAULT_SOURCES_DB,
    DEFAULT_SOURCES_DIR,
    SourceSpec,
    load_sources,
    reject_warehouse_table,
    source_from_mapping,
)

DEFAULT_DB_PATH = DEFAULT_SOURCES_DB

_TABULAR = {".csv", ".tsv"}
_EXCEL = {".xlsx", ".xlsm"}
_ACCEPTED = _TABULAR | _EXCEL
_COLUMN_UNSAFE = re.compile(r"[^a-z0-9_]+")
_COLUMN_REPEAT = re.compile(r"_+")


def load_file(
    file: Path,
    *,
    source_id: str,
    table: str,
    owner_role: str,
    description: str,
    synonyms: tuple[str, ...],
    sheet: str | None = None,
    header_row: int = 1,
    fill_merged: bool = False,
    replace: bool = False,
    as_of_column: str | None = None,
    sources_dir: Path | None = None,
    db_path: Path | None = None,
    routing_path: Path | None = None,
    warehouse_path: Path | None = None,
) -> SourceSpec:
    """Stage ``file`` into the sources database and write its YAML manifest."""
    path = Path(file)
    if not path.is_file():
        raise ConfigError(f"File not found: {path}")
    if header_row < 1:
        raise ConfigError("--header-row must be 1 or greater.")
    suffix = path.suffix.lower()
    _reject_suffix(suffix)
    if suffix in _TABULAR and sheet:
        raise ConfigError("--sheet applies to Excel workbooks only.")
    if suffix in _TABULAR and fill_merged:
        raise ConfigError("--fill-merged applies to Excel workbooks only.")

    if suffix in _TABULAR:
        columns, rows, uncached = _read_tabular(
            path,
            header_row=header_row,
            sep="\t" if suffix == ".tsv" else ",",
        )
    else:
        columns, rows, uncached = _read_excel(
            path,
            sheet=sheet,
            header_row=header_row,
            fill_merged=fill_merged,
        )
    if not columns:
        raise ConfigError(f"{path.name} has no columns to load.")

    manifest_dir = Path(sources_dir) if sources_dir is not None else DEFAULT_SOURCES_DIR
    database = Path(db_path) if db_path is not None else DEFAULT_DB_PATH
    loaded_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    spec = source_from_mapping(
        {
            "source_id": source_id,
            "table": table,
            "owner_role": owner_role,
            "file_name": path.name,
            "sha256": _sha256(path),
            "row_count": len(rows),
            "columns": list(columns),
            "loaded_at": loaded_at,
            "description": description,
            "synonyms": list(synonyms),
            "formula_cells_without_cache": uncached,
            **({"as_of_column": as_of_column.strip()} if as_of_column and as_of_column.strip() else {}),
        },
        filename=f"{source_id.strip()}.yaml",
    )
    previous = _existing(manifest_dir, spec.source_id)
    if previous is not None and previous.sha256 != spec.sha256 and not replace:
        raise ConfigError(
            f"Source {spec.source_id} already exists with a different sha256 "
            f"({previous.sha256} != {spec.sha256}). Pass --replace to overwrite it."
        )
    _reject_duplicate_table(manifest_dir, spec)
    reject_warehouse_table(
        spec.table,
        source_id=spec.source_id,
        routing_path=routing_path,
        warehouse_path=warehouse_path,
    )
    _write_table(
        database,
        spec.table,
        spec.columns,
        rows,
        drop_also=None if previous is None else previous.table,
    )
    _write_manifest(manifest_dir / f"{spec.source_id}.yaml", spec)
    return spec


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m src.sources_load",
        description="Declare a CSV or Excel source. Writing the YAML manifest is the declaration.",
    )
    parser.add_argument("file", type=Path)
    parser.add_argument("--id", required=True, dest="source_id")
    parser.add_argument("--table", required=True)
    parser.add_argument("--owner-role", required=True)
    parser.add_argument("--description", required=True)
    parser.add_argument("--synonym", action="append", default=[], dest="synonyms")
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--fill-merged", action="store_true")
    parser.add_argument("--replace", action="store_true")
    parser.add_argument("--as-of-column", default=None)
    args = parser.parse_args(argv)
    sheet = args.sheet.strip() if isinstance(args.sheet, str) else None
    try:
        spec = load_file(
            args.file,
            source_id=args.source_id,
            table=args.table,
            owner_role=args.owner_role,
            description=args.description,
            synonyms=tuple(args.synonyms),
            sheet=sheet or None,
            header_row=args.header_row,
            fill_merged=args.fill_merged,
            replace=args.replace,
            as_of_column=args.as_of_column,
        )
    except ConfigError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print(f"sha256: {spec.sha256}")
    print(f"row_count: {spec.row_count}")
    print(f"columns: {', '.join(spec.columns)}")
    if spec.formula_cells_without_cache:
        print(
            f"warning: {spec.formula_cells_without_cache} formula cell(s) had no cached value "
            "and were staged as NULL"
        )
    return 0


def _reject_suffix(suffix: str) -> None:
    if suffix == ".xls":
        raise ConfigError(
            "Legacy .xls is not supported: openpyxl reads .xlsx and .xlsm only. "
            "Save the workbook as .xlsx and load that file."
        )
    if suffix not in _ACCEPTED:
        raise ConfigError(
            f"Unsupported file type {suffix or '(none)'}. Accepted types: .csv, .tsv, .xlsx, .xlsm."
        )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _existing(directory: Path, source_id: str) -> SourceSpec | None:
    path = directory / f"{source_id}.yaml"
    if not path.is_file():
        return None
    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    except yaml.YAMLError as exc:
        raise ConfigError(f"Source {path.name} is not valid YAML.") from exc
    return source_from_mapping(raw, filename=path.name)


def _reject_duplicate_table(directory: Path, spec: SourceSpec) -> None:
    if not directory.is_dir():
        return
    for other in load_sources(directory).values():
        if other.table == spec.table and other.source_id != spec.source_id:
            raise ConfigError(f"Duplicate table: {spec.table}")


def _write_manifest(path: Path, spec: SourceSpec) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_id": spec.source_id,
        "table": spec.table,
        "owner_role": spec.owner_role,
        "file_name": spec.file_name,
        "sha256": spec.sha256,
        "row_count": spec.row_count,
        "columns": list(spec.columns),
        "loaded_at": spec.loaded_at,
        "description": spec.description,
        "synonyms": list(spec.synonyms),
        "formula_cells_without_cache": spec.formula_cells_without_cache,
    }
    if spec.as_of_column:
        payload["as_of_column"] = spec.as_of_column
    text = yaml.safe_dump(payload, sort_keys=False, allow_unicode=True)
    temporary = path.with_suffix(".yaml.tmp")
    temporary.write_text(text, encoding="utf-8")
    try:
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _write_table(
    db_path: Path,
    table: str,
    columns: tuple[str, ...],
    rows: list[tuple[object, ...]],
    *,
    drop_also: str | None,
) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    quoted_columns = ", ".join(_quote_ident(name) for name in columns)
    quoted_table = _quote_ident(table)
    placeholders = ", ".join("?" for _ in columns)
    conn = sqlite3.connect(db_path)
    try:
        conn.execute(f"DROP TABLE IF EXISTS {quoted_table}")
        if drop_also and drop_also != table:
            conn.execute(f"DROP TABLE IF EXISTS {_quote_ident(drop_also)}")
        conn.execute(f"CREATE TABLE {quoted_table} ({quoted_columns})")
        if rows:
            conn.executemany(
                f"INSERT INTO {quoted_table} ({quoted_columns}) VALUES ({placeholders})",
                rows,
            )
        conn.commit()
    except sqlite3.Error as exc:
        conn.rollback()
        raise ConfigError(f"Could not write {db_path.name}: {exc}") from exc
    finally:
        conn.close()


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _tabular_width(path: Path, sep: str) -> int:
    """Widest row, so a short title row does not make pandas reject the file."""
    width = 0
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            for row in csv.reader(handle, delimiter=sep):
                width = max(width, len(row))
    except UnicodeDecodeError as exc:
        raise ConfigError(f"Could not read {path.name}: {exc}") from exc
    return width


def _read_tabular(
    path: Path,
    *,
    header_row: int,
    sep: str,
) -> tuple[tuple[str, ...], list[tuple[object, ...]], int]:
    width = _tabular_width(path, sep)
    if width == 0:
        raise ConfigError(f"{path.name} has no rows.")
    try:
        # names= fixes the width up front. A title row with fewer cells is padded
        # instead of being treated as the field count for the whole file.
        frame = pandas.read_csv(
            path,
            sep=sep,
            header=None,
            names=list(range(width)),
            dtype=str,
            na_filter=False,
            skip_blank_lines=False,
            encoding="utf-8-sig",
            keep_default_na=False,
        )
    except pandas.errors.EmptyDataError as exc:
        raise ConfigError(f"{path.name} has no rows.") from exc
    except (pandas.errors.ParserError, UnicodeDecodeError, ValueError) as exc:
        raise ConfigError(f"Could not read {path.name}: {exc}") from exc
    if header_row > len(frame.index):
        raise ConfigError(f"--header-row {header_row} is past the end of {path.name}.")
    raw_header = [_blank(value) for value in frame.iloc[header_row - 1].tolist()]
    body = [
        tuple(_blank(value) for value in frame.iloc[index].tolist())
        for index in range(header_row, len(frame.index))
    ]
    columns, rows = _name_rows(raw_header, _without_trailing_empty(body))
    return columns, rows, 0


def _read_excel(
    path: Path,
    *,
    sheet: str | None,
    header_row: int,
    fill_merged: bool,
) -> tuple[tuple[str, ...], list[tuple[object, ...]], int]:
    try:
        values_book = load_workbook(path, data_only=True, read_only=False)
        formula_book = load_workbook(path, data_only=False, read_only=False)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ConfigError(f"Could not read {path.name}: {exc}") from exc
    try:
        names = list(values_book.sheetnames)
        if not names:
            raise ConfigError(f"{path.name} has no sheets.")
        if sheet is None:
            if len(names) != 1:
                joined = ", ".join(names)
                raise ConfigError(
                    f"Workbook has more than one sheet ({joined}); pass --sheet <name> "
                    "so the wrong sheet is not loaded."
                )
            chosen = names[0]
        elif sheet not in names:
            joined = ", ".join(names)
            raise ConfigError(f"Sheet {sheet!r} was not found. Sheets: {joined}.")
        else:
            chosen = sheet
        return _read_sheet(
            values_book[chosen],
            formula_book[chosen],
            header_row=header_row,
            fill_merged=fill_merged,
            file_name=path.name,
        )
    finally:
        values_book.close()
        formula_book.close()


def _read_sheet(
    values: Worksheet,
    formulas: Worksheet,
    *,
    header_row: int,
    fill_merged: bool,
    file_name: str,
) -> tuple[tuple[str, ...], list[tuple[object, ...]], int]:
    max_row = max(values.max_row or 1, formulas.max_row or 1)
    max_col = max(values.max_column or 1, formulas.max_column or 1)
    for merged in values.merged_cells.ranges:
        max_row = max(max_row, merged.max_row)
        max_col = max(max_col, merged.max_col)
    if header_row > max_row:
        raise ConfigError(f"--header-row {header_row} is past the end of {file_name}.")

    formula_at: dict[tuple[int, int], bool] = {}
    for row in formulas.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            formula_at[(cell.row, cell.column)] = cell.data_type == "f"

    grid: dict[tuple[int, int], object] = {}
    uncached: set[tuple[int, int]] = set()
    for row in values.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            key = (cell.row, cell.column)
            if formula_at.get(key) and cell.value is None:
                grid[key] = None
                uncached.add(key)
            else:
                grid[key] = _plain(cell.value, number_format=cell.number_format)

    for merged in values.merged_cells.ranges:
        origin = grid.get((merged.min_row, merged.min_col))
        for row_index in range(merged.min_row, merged.max_row + 1):
            for col_index in range(merged.min_col, merged.max_col + 1):
                if (row_index, col_index) == (merged.min_row, merged.min_col):
                    continue
                key = (row_index, col_index)
                if fill_merged:
                    grid[key] = origin
                    if origin is not None:
                        uncached.discard(key)
                else:
                    grid[key] = None

    raw_header = [grid.get((header_row, col)) for col in range(1, max_col + 1)]
    body_rows = list(range(header_row + 1, max_row + 1))
    body = [
        tuple(grid.get((row_index, col)) for col in range(1, max_col + 1))
        for row_index in body_rows
    ]
    kept_columns = _kept_column_indexes(raw_header, body, body_rows, uncached)
    if not kept_columns:
        raise ConfigError(f"{file_name} has no columns to load.")
    excel_columns = [index + 1 for index in kept_columns]
    header = [raw_header[index] for index in kept_columns]
    narrowed = [tuple(record[index] for index in kept_columns) for record in body]
    kept_rows = _without_trailing_empty_formulas(narrowed, body_rows, excel_columns, uncached)
    columns, rows = _name_rows(header, kept_rows)
    staged_excel_rows = body_rows[: len(rows)]
    count = sum(
        (excel_row, excel_col) in uncached and rows[offset][col_offset] is None
        for offset, excel_row in enumerate(staged_excel_rows)
        for col_offset, excel_col in enumerate(excel_columns)
    )
    return columns, rows, count


def _kept_column_indexes(
    header: list[object],
    rows: list[tuple[object, ...]],
    excel_rows: list[int],
    uncached: set[tuple[int, int]],
) -> list[int]:
    width = len(header)
    last = width
    while last > 0:
        index = last - 1
        excel_col = index + 1
        header_empty = header[index] is None
        data_empty = all(index >= len(record) or record[index] is None for record in rows)
        has_formula = any((excel_row, excel_col) in uncached for excel_row in excel_rows)
        if header_empty and data_empty and not has_formula:
            last -= 1
            continue
        break
    return list(range(last))


def _without_trailing_empty_formulas(
    rows: list[tuple[object, ...]],
    excel_rows: list[int],
    excel_columns: list[int],
    uncached: set[tuple[int, int]],
) -> list[tuple[object, ...]]:
    end = len(rows)
    while end > 0:
        record = rows[end - 1]
        excel_row = excel_rows[end - 1]
        if any(value is not None for value in record):
            break
        if any((excel_row, excel_col) in uncached for excel_col in excel_columns):
            break
        end -= 1
    return rows[:end]


def _name_rows(
    raw_header: list[object],
    rows: list[tuple[object, ...]],
) -> tuple[tuple[str, ...], list[tuple[object, ...]]]:
    width = len(raw_header)
    padded = [tuple(list(record) + [None] * (width - len(record)))[:width] for record in rows]
    return _normalise_headers(raw_header), padded


def _without_trailing_empty(rows: list[tuple[object, ...]]) -> list[tuple[object, ...]]:
    end = len(rows)
    while end > 0 and all(value is None for value in rows[end - 1]):
        end -= 1
    return rows[:end]


def _normalise_headers(raw_header: list[object]) -> tuple[str, ...]:
    seen: dict[str, int] = {}
    names: list[str] = []
    for index, raw in enumerate(raw_header):
        base = _normalise_header(raw, index)
        seen[base] = seen.get(base, 0) + 1
        count = seen[base]
        names.append(base if count == 1 else f"{base}_{count}")
    return tuple(names)


def _normalise_header(raw: object, index: int) -> str:
    text = "" if raw is None else str(raw).strip().lower()
    cleaned = _COLUMN_REPEAT.sub("_", _COLUMN_UNSAFE.sub("_", text)).strip("_")
    if not cleaned:
        return f"column_{index + 1}"
    if cleaned[0].isdigit():
        return f"column_{cleaned}"
    return cleaned


def _blank(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, float) and value != value:
        return None
    if isinstance(value, str):
        return None if value == "" else value
    return value


def _plain(value: object, *, number_format: str | None = None) -> object:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, datetime):
        has_clock = bool(value.hour or value.minute or value.second or value.microsecond)
        if has_clock or (number_format and _format_has_time(number_format)):
            return value.isoformat(sep=" ")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, (int, float, str)):
        if isinstance(value, str) and value == "":
            return None
        return value
    return str(value)


def _format_has_time(number_format: str) -> bool:
    stripped = re.sub(r'"[^"]*"', "", number_format.lower())
    return re.search(r"h|s|am/pm", stripped) is not None


if __name__ == "__main__":
    raise SystemExit(main())
